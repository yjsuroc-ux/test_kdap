"""파일 읽기: 온나라 접수 목록(엑셀/CSV/HTML형 xls)과 본문(HWP 5.0 / HWPX / TXT)."""
import csv
import datetime as dt
import html
import io
import re
import struct
import zipfile
import zlib
from html.parser import HTMLParser

EXCEL_EXTS = (".xlsx", ".xlsm", ".xls", ".csv")
BODY_EXTS = (".hwp", ".hwpx", ".txt")


class UnreadableFile(Exception):
    pass


# ================================================================ HWP 5.0

HWPTAG_PARA_TEXT = 67  # HWPTAG_BEGIN(16) + 51

# 확장/인라인 컨트롤은 8 글자(16바이트)를 차지한다. 나머지 0~31은 1 글자.
_CTRL_8 = {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23}


def _iter_records(buf):
    pos, n = 0, len(buf)
    while pos + 4 <= n:
        (h,) = struct.unpack_from("<I", buf, pos)
        pos += 4
        tag, size = h & 0x3FF, (h >> 20) & 0xFFF
        if size == 0xFFF:
            if pos + 4 > n:
                break
            (size,) = struct.unpack_from("<I", buf, pos)
            pos += 4
        yield tag, buf[pos:pos + size]
        pos += size


def _para_text(data):
    n = len(data) // 2
    units = struct.unpack("<%dH" % n, data[:n * 2])
    out, i = [], 0
    while i < n:
        c = units[i]
        if c < 32:
            if c == 9:
                out.append(9)
            elif c in (10, 13):
                out.append(10)
            i += 8 if c in _CTRL_8 else 1
        else:
            out.append(c)
            i += 1
    return struct.pack("<%dH" % len(out), *out).decode("utf-16le", errors="replace")


def extract_hwp(data):
    import olefile

    if not olefile.isOleFile(data=data):
        raise UnreadableFile("HWP 5.0 형식이 아님")
    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        if not ole.exists("FileHeader"):
            raise UnreadableFile("HWP 파일 헤더 없음")
        header = ole.openstream("FileHeader").read()
        if not header.startswith(b"HWP Document File"):
            raise UnreadableFile("HWP 파일 서명이 아님")
        (props,) = struct.unpack_from("<I", header, 36)
        if props & 0x02:
            raise UnreadableFile("암호가 걸린 HWP")
        if props & 0x04:
            raise UnreadableFile("배포용 HWP(암호화)")
        compressed = bool(props & 0x01)
        sections = []
        for entry in ole.listdir(streams=True, storages=False):
            if len(entry) == 2 and entry[0] == "BodyText" and entry[1].startswith("Section"):
                try:
                    sections.append((int(entry[1][7:]), entry))
                except ValueError:
                    pass
        if not sections:
            raise UnreadableFile("본문(BodyText) 없음")
        parts = []
        for _, entry in sorted(sections):
            raw = ole.openstream(entry).read()
            if compressed:
                raw = zlib.decompressobj(-15).decompress(raw)
            for tag, rec in _iter_records(raw):
                if tag == HWPTAG_PARA_TEXT:
                    t = _para_text(rec)
                    parts.append(t if t.endswith("\n") else t + "\n")
        return "".join(parts)
    finally:
        ole.close()


# ================================================================ HWPX

def extract_hwpx(data):
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise UnreadableFile("HWPX 형식이 아님")
    names = sorted(
        (n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n)),
        key=lambda n: int(re.findall(r"\d+", n)[-1]),
    )
    if not names:
        raise UnreadableFile("HWPX 본문 없음(배포용일 수 있음)")
    out = []
    for n in names:
        xml = z.read(n).decode("utf-8", errors="replace")
        xml = re.sub(r"<(\w+:)?tab\b[^>]*/>", "\t", xml)
        xml = re.sub(r"<(\w+:)?lineBreak\b[^>]*/>", "\n", xml)
        xml = re.sub(r"</(\w+:)?p>", "\n", xml)
        xml = re.sub(r"<[^>]+>", "", xml)
        out.append(html.unescape(xml))
    return "".join(out).strip("\n")


def _decode_text(data):
    for enc in ("utf-8-sig", "cp949", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_body(filename, data):
    name = filename.lower()
    if name.endswith(".hwp"):
        return extract_hwp(data)
    if name.endswith(".hwpx"):
        return extract_hwpx(data)
    if name.endswith(".txt"):
        return _decode_text(data)
    raise UnreadableFile("지원하지 않는 형식")


# ================================================================ 엑셀/CSV

class _TableParser(HTMLParser):
    """온나라 등에서 .xls 확장자로 내려주는 HTML 표를 읽는다."""

    def __init__(self):
        super().__init__()
        self.rows, self._row, self._cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(re.sub(r"\s+", " ", "".join(self._cell)).strip())
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, d):
        if self._cell is not None:
            self._cell.append(d)


def _read_sheets(filename, data):
    """→ [(시트 이름, 행 목록)]. 통합문서의 모든 시트를 읽는다."""
    name = filename.lower()
    head = data[:2048].lstrip().lower()
    if head.startswith(b"<") or b"<table" in head or b"<html" in head:
        p = _TableParser()
        p.feed(_decode_text(data))
        return [("", p.rows)]
    if name.endswith(".csv"):
        return [("", list(csv.reader(io.StringIO(_decode_text(data)))))]
    if name.endswith(".xls"):
        try:
            import xlrd
        except ImportError:
            raise UnreadableFile(".xls 파일을 읽으려면 xlrd 패키지가 필요합니다. .xlsx로 저장해 올려 주세요.")
        book = xlrd.open_workbook(file_contents=data)
        sheets = []
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                    else:
                        row.append(cell.value)
                rows.append(row)
            sheets.append((sh.name, rows))
        return sheets
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()


# 온나라 접수문서 목록의 열 이름 후보. 위에서부터 정확히 같은 이름을 먼저 찾는다.
HEADER_SYNONYMS = {
    "doc_no": ["문서번호", "생산문서번호", "시행번호", "발신문서번호"],
    "receipt_no": ["접수번호", "접수문서번호"],
    "title": ["제목", "문서제목"],
    "sender": ["발신기관", "발신처", "생산기관", "발신기관명"],
    "sender_dept": ["발신부서", "생산부서", "발신부서명"],
    "produced": ["생산일자", "생산일", "생산일시", "시행일자", "발신일자"],
    "received": ["접수일시", "접수일", "접수일자", "접수일자시"],
    "assignee": ["처리자", "담당자", "처리담당자", "배부자", "접수자"],
    "status": ["결재상태", "처리상태", "문서상태", "상태"],
}


def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


_ALL_SYN = {re.sub(r"\s", "", s) for v in HEADER_SYNONYMS.values() for s in v}


def _header_score(r):
    filled = sum(1 for c in r if c)
    known = sum(1 for c in r if re.sub(r"\s", "", c) in _ALL_SYN)
    return known * 10 + filled


def read_list(filename, data):
    """→ (headers, rows, sheet). 여러 시트 중 접수 목록으로 보이는 시트를 고르고,
    앞쪽의 제목 줄 등을 건너뛰어 머리글 행을 자동으로 찾는다."""
    best = None
    for sheet, raw in _read_sheets(filename, data):
        rows = [[cell_str(v) for v in r] for r in raw]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        hi = max(range(min(20, len(rows))), key=lambda i: (_header_score(rows[i]), -i))
        score = (_header_score(rows[hi]), len(rows))
        if best is None or score > best[0]:
            best = (score, sheet, rows, hi)
    if best is None:
        raise UnreadableFile("엑셀에 내용이 없습니다")
    _, sheet, rows, hi = best
    headers, seen = [], {}
    for i, h in enumerate(rows[hi]):
        h = h or f"열{i + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h}({seen[h]})"
        else:
            seen[h] = 1
        headers.append(h)
    body = []
    for r in rows[hi + 1:]:
        r = (r + [""] * len(headers))[: len(headers)]
        if any(r):
            body.append(r)
    return headers, body, sheet


def guess_mapping(headers):
    """열 이름으로 매핑을 추정한다. 같은 열을 두 항목에 쓰지 않는다."""
    mapping, used = {}, set()
    norm = [re.sub(r"\s", "", h) for h in headers]
    for field, syns in HEADER_SYNONYMS.items():  # 1차: 이름이 정확히 같은 열
        for s in syns:
            if s in norm and norm.index(s) not in used:
                mapping[field] = headers[norm.index(s)]
                used.add(norm.index(s))
                break
    for field, syns in HEADER_SYNONYMS.items():  # 2차: 이름에 후보가 들어 있는 열
        if field in mapping:
            continue
        for i, h in enumerate(norm):
            if i not in used and any(s in h for s in syns):
                mapping[field] = headers[i]
                used.add(i)
                break
    return mapping



def parse_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    m = re.search(r"((?:19|20)\d{2})\s*[.\-/년]\s*(\d{1,2})\s*[.\-/월]\s*(\d{1,2})", s)
    if not m:
        m = re.match(r"((?:19|20)\d{2})(\d{2})(\d{2})", s)
    if m:
        try:
            return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    if re.fullmatch(r"\d{5}(\.\d+)?", s):  # 엑셀 일련번호
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(float(s)))
    return None
